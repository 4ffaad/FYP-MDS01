#!/usr/bin/env python3
"""Build the Google Docs import-ready report from the Markdown master."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "final-report.md"
OUTPUT = ROOT / "MDS01_Final_Project_Report.docx"

NAVY = "19334D"
BLUE = "245F8C"
PALE_BLUE = "EAF2F8"
MID_GREY = "586B7D"
CODE_BG = "F2F5F7"


def set_cell_shading(cell, fill: str) -> None:
    """Apply a solid background fill to one table cell."""
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_margins(cell, top=65, start=85, bottom=65, end=85) -> None:
    """Set compact cell padding in twips for report tables."""
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    """Tell Word to repeat the table's first row after page breaks."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def prevent_row_split(row) -> None:
    """Keep each short report-table row together when it fits on a page."""
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph) -> None:
    """Insert a live Word page-number field."""
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend((begin, instruction, separate, value, end))
    field_run.font.size = Pt(8)


def set_section_geometry(section, *, landscape: bool) -> None:
    """Set A4 page geometry and practical margins for portrait/RTM pages."""
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.68)
        section.right_margin = Inches(0.68)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.68)


def add_header_footer(section, *, first=False) -> None:
    """Add restrained running identification and page numbering."""
    if first:
        section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.style = "Normal"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("MDS01  |  FINAL PROJECT REPORT")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)

    footer = section.footer.paragraphs[0]
    footer.style = "Normal"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    add_page_field(footer)


def clean_inline(text: str) -> str:
    """Remove Markdown emphasis and link syntax while keeping readable labels."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*(?!\*)", r"\1", text)
    return text


def add_inline(paragraph, text: str, *, code_default=False) -> None:
    """Render basic Markdown emphasis and inline-code runs into Word."""
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+\*(?!\*))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(clean_inline(text[cursor:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(clean_inline(token[2:-2]))
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(clean_inline(text[cursor:]))
    if code_default:
        for run in paragraph.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(7.2)


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    """Split a Markdown pipe table, skipping the separator row."""
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell or "-") for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_markdown_table(doc: Document, rows: list[list[str]], *, landscape: bool) -> None:
    """Create a readable Word table and style its header and evidence text."""
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = False
    available = 10.55 if landscape else 6.91
    if columns == 5 and landscape:
        widths = [0.62, 2.15, 1.08, 3.12, 3.58]
    elif columns == 3:
        widths = [1.7, 2.55, available - 4.25]
    else:
        widths = [available / columns] * columns
    for row_index, source_row in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        for col_index in range(columns):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            text = source_row[col_index] if col_index < len(source_row) else ""
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(7.0 if columns == 5 else 8.0)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F7FA")
    repeat_table_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)


def add_code_block(doc: Document, lines: list[str]) -> None:
    """Insert source code with preserved indentation and a readable print size."""
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.10)
        paragraph.paragraph_format.right_indent = Inches(0.05)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.92
        ppr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), CODE_BG)
        ppr.append(shading)
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(7.2)
        run.font.color.rgb = RGBColor.from_string("253746")


def add_image(doc: Document, alt: str, relative_path: str, source_path: Path, *, landscape: bool) -> None:
    """Insert one centered figure, retaining descriptive alternative text."""
    path = (source_path.parent / relative_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"image referenced in report is missing: {path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    if landscape:
        width = Inches(10.45)
    elif path.name == "vsvig-gradcam-evidence.png":
        width = Inches(6.25)
    else:
        width = Inches(6.85)
    inline_shape = paragraph.add_run().add_picture(str(path), width=width)
    inline_shape._inline.docPr.set("descr", alt)


def set_document_styles(doc: Document) -> None:
    """Set a consistent accessible type scale and restrained report palette."""
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("243447")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 25, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 10.5, BLUE)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(10)
        styles[name].paragraph_format.space_after = Pt(2)


def add_cover(doc: Document, markdown: str) -> None:
    """Create a clean report title page from the Markdown front matter."""
    lines = markdown.splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), "MDS01 Final Project Report")
    cover = doc.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Inches(1.6)
    cover.paragraph_format.space_after = Pt(12)
    add_inline(cover, title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("MDS01 Final Project Report")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    metadata = [
        "Daffa Fariq Prasetyo (34076492)",
        "Syed Asad Ali Shah (34726918)",
        "Jeevinaash Vijayenthran (34828532)",
        "Shan Pasidu Dissanayake (31098703)",
        "Supervisor: Dr Jessie Leong Shu Min",
        "FIT3162 / FIT3164 / FIT3189",
        "27 October 2026",
        "Main body: approximately 2,500 words (maximum 3,500)",
    ]
    for item in metadata:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(item)
        r.font.size = Pt(11 if item.startswith("Main body") else 10.5)
        if item.startswith("Main body"):
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(MID_GREY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("Draft for team and supervisor review")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)
    doc.add_page_break()


def add_section(doc: Document, landscape: bool):
    """Begin a new page section and switch orientation where required."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section, landscape=landscape)
    # New sections inherit the single running header/footer. Appending new
    # runs here would duplicate them because the headers are linked by default.
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    return section


def build() -> Path:
    """Transform the report Markdown into a formatted Word document."""
    markdown = SOURCE.read_text()
    doc = Document()
    set_section_geometry(doc.sections[0], landscape=False)
    add_header_footer(doc.sections[0], first=True)
    set_document_styles(doc)
    doc.core_properties.title = "Secure EEG VEEG Clinical Review and Experimental Video Analysis"
    doc.core_properties.subject = "MDS01 final project report"
    doc.core_properties.author = "MDS01 Project Team"
    add_cover(doc, markdown)

    start = markdown.index("### Authorship statement")
    lines = markdown[start:].splitlines()
    landscape = False
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines)
            paragraph = doc.add_paragraph()
            if text.startswith("**Figure "):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = False
            add_inline(paragraph, text)
            paragraph_lines = []

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if in_code:
            if stripped.startswith("```"):
                add_code_block(doc, [code_line for code_line in code_lines if code_line.strip()])
                code_lines = []
                in_code = False
            else:
                code_lines.append(line.rstrip())
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            in_code = True
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(doc, parse_table_rows(table_lines), landscape=landscape)
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            flush_paragraph()
            add_image(doc, image_match.group(1), image_match.group(2), SOURCE, landscape=landscape)
            index += 1
            continue
        if stripped.startswith("## ") or stripped.startswith("### "):
            flush_paragraph()
            level = 1 if stripped.startswith("## ") else 2
            title = stripped[level + 1:].strip()
            if title.startswith("1. Introduction") or title.startswith("7. References"):
                doc.add_page_break()
            if title.startswith("Appendix A."):
                add_section(doc, landscape=True)
                landscape = True
            elif title.startswith("Appendix B."):
                add_section(doc, landscape=False)
                landscape = False
            elif title.startswith("Appendix C.") or title.startswith("Appendix D."):
                doc.add_page_break()
            heading = doc.add_paragraph(style=f"Heading {level}")
            add_inline(heading, title)
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            number, item_text = re.match(r"^(\d+)\.\s+(.*)$", stripped).groups()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.23)
            paragraph.paragraph_format.first_line_indent = Inches(-0.23)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(f"{number}.\t")
            add_inline(paragraph, item_text)
            index += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1

    flush_paragraph()
    if in_code:
        add_code_block(doc, [code_line for code_line in code_lines if code_line.strip()])
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
